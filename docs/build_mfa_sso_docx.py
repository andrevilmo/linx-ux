#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Telas estilo Portal + DOCX (frontend + parâmetros SSO da empresa)."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
import qrcode

IMG = Path("/tmp/mfa-sso-docx/img")
ART = Path("/opt/cursor/artifacts")
DOCS = Path("/workspace/docs")
DOCX_NAME = "Linx-UX-MFA-SSO-Portal.docx"

FONT = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONTB = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
FONTI = "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf"

PURPLE = (91, 46, 144)
PURPLE_D = (62, 28, 102)
WHITE = (255, 255, 255)
GRAY = (117, 117, 117)
LINE = (207, 207, 207)
NAVY_T = (11, 37, 69)
SLATE_T = (51, 65, 85)
BG_L = (32, 18, 58)
ERR = (198, 40, 40)

NAVY = RGBColor(0x0B, 0x25, 0x45)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
PURPLE_RGB = RGBColor(0x5B, 0x2E, 0x90)


def F(path, size):
    return ImageFont.truetype(path, size)


def rr(d, box, r, fill, outline=None, w=2):
    d.rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=w)


def portal_frame(title_right_extra=None):
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), WHITE)
    d = ImageDraw.Draw(img)
    # left fashion panel
    d.rectangle((0, 0, 300, H), fill=BG_L)
    for i in range(20):
        d.rectangle((0, i * 44, 300, i * 44 + 22), fill=(42, 24, 72))
    d.rectangle((0, 0, 300, H), outline=None)
    d.text((28, 36), "LINX", font=F(FONTB, 34), fill=WHITE)
    d.text((28, 92), "SOFTWARE QUE", font=F(FONT, 16), fill=(230, 220, 245))
    d.text((28, 116), "MOVE O VAREJO", font=F(FONT, 16), fill=(230, 220, 245))
    d.text((16, H - 36), "© Linx - Todos os direitos reservados.", font=F(FONT, 12), fill=(200, 190, 210))
    # right
    d.text((360, 70), "linx", font=F(FONTB, 56), fill=(30, 30, 30))
    tw = d.textlength("linx", font=F(FONTB, 56))
    d.text((360 + tw, 70), "ux", font=F(FONTB, 56), fill=PURPLE)
    d.text((W - 360, H - 28), "© Linx - Todos os direitos reservados.", font=F(FONT, 13), fill=GRAY)
    return img, d


def underline_field(d, x, y, w, label, value, readonly=False):
    d.text((x, y), label.upper(), font=F(FONT, 15), fill=GRAY)
    d.text((x, y + 32), value, font=F(FONT, 22), fill=(40, 40, 40) if value else (180, 180, 180))
    d.line((x, y + 68, x + w, y + 68), fill=PURPLE if readonly else LINE, width=2)


def purple_btn(d, box, text, ghost=False):
    if ghost:
        rr(d, box, 2, (238, 238, 238), PURPLE, 1)
        d.text(((box[0] + box[2]) / 2, (box[1] + box[3]) / 2), text.upper(), font=F(FONT, 16), fill=PURPLE, anchor="mm")
    else:
        rr(d, box, 2, PURPLE)
        d.text(((box[0] + box[2]) / 2, (box[1] + box[3]) / 2), text.upper(), font=F(FONT, 16), fill=WHITE, anchor="mm")


def shot_identificar():
    img, d = portal_frame()
    underline_field(d, 360, 240, 680, "usuário", "")
    purple_btn(d, (360, 360, 1040, 410), "continuar")
    p = IMG / "tela-identificar.png"
    img.save(p, "PNG")
    return p


def shot_senha():
    img, d = portal_frame()
    underline_field(d, 360, 210, 680, "usuário", "joao.silva", True)
    d.text((360, 290), "alterar usuário", font=F(FONT, 16), fill=PURPLE)
    underline_field(d, 360, 330, 680, "senha", "••••••••")
    d.rectangle((980, 412, 996, 428), outline=PURPLE, width=2)
    d.rectangle((983, 415, 993, 425), fill=PURPLE)
    d.text((1004, 412), "mantenha-me conectado", font=F(FONT, 15), fill=GRAY)
    purple_btn(d, (360, 460, 740, 508), "listar ambientes", ghost=True)
    purple_btn(d, (760, 460, 1040, 508), "entrar")
    d.text((700, 540), "ou", font=F(FONT, 16), fill=(160, 160, 160), anchor="mm")
    purple_btn(d, (360, 568, 1040, 620), "entrar com Microsoft")
    d.text((360, 780), "Esqueci minha senha", font=F(FONT, 16), fill=PURPLE)
    p = IMG / "tela-senha-sso.png"
    img.save(p, "PNG")
    return p


def shot_microsoft():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), (242, 242, 242))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 52), fill=WHITE)
    d.line((0, 52, W, 52), fill=(225, 225, 225), width=1)
    # ms squares
    d.rectangle((24, 14, 34, 24), fill=(242, 80, 34))
    d.rectangle((36, 14, 46, 24), fill=(127, 186, 0))
    d.rectangle((24, 26, 34, 36), fill=(0, 164, 239))
    d.rectangle((36, 26, 46, 36), fill=(255, 185, 0))
    d.text((58, 26), "login.microsoftonline.com", font=F(FONT, 16), fill=(97, 97, 97), anchor="lm")
    # card
    rr(d, (500, 160, 940, 620), 2, WHITE, (220, 220, 220), 1)
    d.rectangle((532, 196, 542, 206), fill=(242, 80, 34))
    d.rectangle((544, 196, 554, 206), fill=(127, 186, 0))
    d.rectangle((532, 208, 542, 218), fill=(0, 164, 239))
    d.rectangle((544, 208, 554, 218), fill=(255, 185, 0))
    d.text((532, 250), "Entrar", font=F(FONTB, 32), fill=(27, 27, 27))
    d.text((532, 300), "joao.silva@empresa.com", font=F(FONT, 18), fill=(97, 97, 97))
    d.text((532, 380), "Senha", font=F(FONT, 16), fill=(141, 141, 141))
    d.line((532, 420, 908, 420), fill=(138, 136, 134), width=1)
    rr(d, (760, 470, 908, 514), 2, (0, 103, 184))
    d.text((834, 492), "Avançar", font=F(FONT, 16), fill=WHITE, anchor="mm")
    p = IMG / "tela-microsoft.png"
    img.save(p, "PNG")
    return p


def shot_mfa(qr=False):
    img, d = portal_frame()
    d.text((360, 200), "Verificação em duas etapas", font=F(FONT, 26), fill=(80, 80, 80))
    y = 250
    if qr:
        d.text((360, y), "Escaneie o QR Code no aplicativo de autenticação.", font=F(FONT, 18), fill=GRAY)
        d.text((360, y + 32), "ACME  ·  joao.silva", font=F(FONT, 16), fill=GRAY)
        q = qrcode.make("otpauth://totp/ACME:joao.silva?secret=DEMOSECRET&issuer=ACME")
        q = q.resize((148, 148))
        img.paste(q, (360, y + 70))
        y = y + 240
    else:
        d.text((360, y), "ACME  ·  joao.silva", font=F(FONT, 18), fill=GRAY)
        y = y + 40
    d.text((360, y), "Para continuar é necessário inserir o código gerado", font=F(FONT, 18), fill=GRAY)
    d.text((360, y + 28), "pelo aplicativo de autenticação", font=F(FONT, 18), fill=GRAY)
    d.text((360, y + 80), "CÓDIGO DE 6 DÍGITOS", font=F(FONT, 14), fill=GRAY)
    rr(d, (360, y + 110, 700, y + 168), 2, (250, 248, 255), PURPLE, 2)
    d.text((530, y + 139), "482915" if not qr else "000000", font=F(FONTB, 28), fill=PURPLE if not qr else (180, 180, 180), anchor="mm")
    purple_btn(d, (360, y + 190, 560, y + 238), "continuar")
    d.text((360, 800), "voltar aos ambientes", font=F(FONT, 14), fill=GRAY)
    name = "tela-mfa-qr.png" if qr else "tela-mfa-codigo.png"
    p = IMG / name
    img.save(p, "PNG")
    return p


def shot_ambientes():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), (26, 16, 48))
    d = ImageDraw.Draw(img)
    d.text((48, 36), "LINX", font=F(FONTB, 28), fill=WHITE)
    d.text((48, 120), "Ambientes disponíveis:", font=F(FONT, 26), fill=WHITE)
    d.rectangle((48, 180, 1390, 248), fill=PURPLE)
    for i, h in enumerate(["Ambiente", "Empresa", "Grupo econômico", ""]):
        d.text((70 + i * 330, 214), h, font=F(FONTB, 16), fill=WHITE, anchor="lm")
    rows = [("Produção", "ACME Varejo Ltda", "ACME"), ("Homologação", "ACME Varejo Ltda", "ACME")]
    for r, row in enumerate(rows):
        y = 248 + r * 64
        d.rectangle((48, y, 1390, y + 64), fill=WHITE if r == 0 else (248, 248, 252))
        for i, val in enumerate(row):
            d.text((70 + i * 330, y + 32), val, font=F(FONT, 18), fill=(60, 60, 60), anchor="lm")
        d.text((1220, y + 32), "entrar", font=F(FONT, 18), fill=PURPLE, anchor="lm")
    p = IMG / "tela-ambientes.png"
    img.save(p, "PNG")
    return p


def shot_cadastro():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), (245, 246, 248))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 56), fill=PURPLE)
    d.text((32, 28), "Linx UX  ·  Cadastro de usuário de autenticação", font=F(FONT, 20), fill=WHITE, anchor="lm")
    rr(d, (80, 100, 1360, 780), 8, WHITE, (220, 220, 220), 1)
    d.text((112, 130), "Dados de autenticação", font=F(FONTB, 22), fill=PURPLE)
    fields = [("Nome de autenticação", "joao.silva"), ("E-mail", "joao.silva@empresa.com"), ("Nome", "João Silva")]
    y = 190
    for lab, val in fields:
        d.text((112, y), lab, font=F(FONT, 15), fill=(90, 90, 90))
        rr(d, (380, y - 8, 1280, y + 36), 3, WHITE, LINE, 1)
        d.text((394, y + 14), val, font=F(FONT, 17), fill=(40, 40, 40), anchor="lm")
        y += 64
    y += 10
    for label, on in (("Utiliza SSO", True), ("Utiliza MFA", True), ("Usuário de serviço", False)):
        rr(d, (112, y, 136, y + 24), 3, PURPLE if on else WHITE, PURPLE, 2)
        if on:
            d.line((118, y + 13, 124, y + 19), fill=WHITE, width=2)
            d.line((124, y + 19, 132, y + 7), fill=WHITE, width=2)
        d.text((148, y + 12), label, font=F(FONT, 17), fill=(50, 50, 50), anchor="lm")
        y += 40
    d.text((112, y + 8), "O e-mail é enviado à Microsoft como dica de conta (login_hint).", font=F(FONTI, 15), fill=GRAY)
    d.text((112, y + 32), "O vínculo Azure é gravado no primeiro SSO bem-sucedido.", font=F(FONTI, 15), fill=GRAY)
    purple_btn(d, (112, 680, 300, 728), "Revogar MFA", ghost=True)
    purple_btn(d, (320, 680, 520, 728), "Revogar SSO", ghost=True)
    purple_btn(d, (540, 680, 700, 728), "Salvar")
    p = IMG / "tela-cadastro.png"
    img.save(p, "PNG")
    return p


def shot_erro():
    img, d = portal_frame()
    d.text((360, 200), "Conta Microsoft diferente da vinculada a este usuário.", font=F(FONT, 18), fill=ERR)
    d.text((360, 228), "Use a conta Azure já associada ou peça para revogar o SSO no cadastro.", font=F(FONT, 16), fill=ERR)
    underline_field(d, 360, 300, 680, "usuário", "joao.silva")
    purple_btn(d, (360, 420, 1040, 470), "continuar")
    p = IMG / "tela-vinculo-erro.png"
    img.save(p, "PNG")
    return p


def shot_jornada():
    W, H = 1600, 420
    img = Image.new("RGB", (W, H), (244, 247, 251))
    d = ImageDraw.Draw(img)
    d.text((60, 24), "Jornada no Portal (o que a pessoa vê)", font=F(FONTB, 28), fill=NAVY_T)
    steps = [
        ("1", "Usuário", "CONTINUAR", (13, 115, 119)),
        ("2", "Senha ou", "Microsoft", PURPLE),
        ("3", "Ambiente", "se houver vários", (27, 79, 138)),
        ("4", "MFA Linx", "QR ou 6 dígitos", (224, 122, 61)),
        ("5", "Application", "produto aberto", (46, 125, 79)),
    ]
    for i, (n, a, b, c) in enumerate(steps):
        x = 50 + i * 310
        rr(d, (x, 110, x + 270, 360), 18, WHITE, c, 4)
        d.ellipse((x + 20, 132, x + 78, 190), fill=c)
        d.text((x + 49, 161), n, font=F(FONTB, 26), fill=WHITE, anchor="mm")
        d.text((x + 24, 220), a, font=F(FONTB, 24), fill=NAVY_T)
        d.text((x + 24, 268), b, font=F(FONT, 20), fill=SLATE_T)
        if i < 4:
            d.polygon([(x + 278, 235), (x + 302, 225), (x + 302, 245)], fill=c)
    p = IMG / "jornada-portal.png"
    img.save(p, "PNG")
    return p


def shot_jornada_sso():
    W, H = 1600, 520
    img = Image.new("RGB", (W, H), (244, 247, 251))
    d = ImageDraw.Draw(img)
    d.text((60, 20), "Processo completo de SSO (o que acontece por baixo das telas)", font=F(FONTB, 26), fill=NAVY_T)
    steps = [
        ("1", "CONTINUAR", "IDENT", (13, 115, 119)),
        ("2", "Microsoft", "START + AZURE", PURPLE),
        ("3", "Vínculo", "FIRST ou LINK", (27, 79, 138)),
        ("4", "Sessão", "Login SSO", (46, 125, 79)),
        ("5", "Ambiente", "+ MFA Linx", (224, 122, 61)),
        ("6", "Revogar", "REV no cadastro", (153, 27, 91)),
    ]
    for i, (n, a, b, c) in enumerate(steps):
        x = 36 + i * 260
        rr(d, (x, 90, x + 236, 340), 16, WHITE, c, 4)
        d.ellipse((x + 16, 110, x + 68, 162), fill=c)
        d.text((x + 42, 136), n, font=F(FONTB, 24), fill=WHITE, anchor="mm")
        d.text((x + 16, 188), a, font=F(FONTB, 20), fill=NAVY_T)
        d.text((x + 16, 230), b, font=F(FONT, 16), fill=SLATE_T)
        if i < 5:
            d.polygon([(x + 242, 210), (x + 256, 200), (x + 256, 220)], fill=c)
    d.text((60, 380), "1–5 = login no Portal.  6 = Revogar SSO no Application (apaga só o vínculo; o próximo Microsoft grava FIRST de novo).", font=F(FONT, 18), fill=SLATE_T)
    d.text((60, 420), "Mismatch de conta Azure (OID diferente) interrompe no passo 3, grava SSOF-LINK e volta ao CONTINUAR.", font=F(FONT, 18), fill=ERR)
    d.text((60, 456), "Canal TCS_LOG_ACESSO_AUTH = PortalSSO.  I = passo  ·  F = falha sem lockout  ·  S = Login SSO efetuado.", font=F(FONT, 16), fill=GRAY)
    p = IMG / "jornada-sso.png"
    img.save(p, "PNG")
    return p


def shot_revogar_confirm():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), (245, 246, 248))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 56), fill=PURPLE)
    d.text((32, 28), "Linx UX  ·  Cadastro de usuário de autenticação", font=F(FONT, 20), fill=WHITE, anchor="lm")
    rr(d, (80, 100, 1360, 780), 8, WHITE, (220, 220, 220), 1)
    d.text((112, 130), "Dados de autenticação", font=F(FONTB, 22), fill=PURPLE)
    d.text((112, 190), "Nome de autenticação", font=F(FONT, 15), fill=(90, 90, 90))
    rr(d, (380, 182, 1280, 226), 3, WHITE, LINE, 1)
    d.text((394, 204), "joao.silva", font=F(FONT, 17), fill=(40, 40, 40), anchor="lm")
    purple_btn(d, (112, 680, 300, 728), "Revogar MFA", ghost=True)
    purple_btn(d, (320, 680, 520, 728), "Revogar SSO", ghost=True)
    # dim overlay
    overlay = Image.new("RGBA", (W, H), (20, 16, 32, 110))
    img = img.convert("RGBA")
    img = Image.alpha_composite(img, overlay)
    d = ImageDraw.Draw(img)
    rr(d, (380, 250, 1060, 560), 8, WHITE, (200, 200, 200), 1)
    d.text((720, 292), "Revogar SSO", font=F(FONTB, 24), fill=PURPLE, anchor="mm")
    d.text((420, 340), "Revogar o SSO deste usuário? Remove só o vínculo", font=F(FONT, 17), fill=(50, 50, 50))
    d.text((420, 368), "da conta Microsoft. O próximo login SSO gravará", font=F(FONT, 17), fill=(50, 50, 50))
    d.text((420, 396), "um novo OID/UPN.", font=F(FONT, 17), fill=(50, 50, 50))
    purple_btn(d, (520, 470, 700, 518), "Não", ghost=True)
    purple_btn(d, (740, 470, 920, 518), "Sim")
    img = img.convert("RGB")
    p = IMG / "tela-revogar-sso.png"
    img.save(p, "PNG")
    return p


def shot_revogar_ok():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), (245, 246, 248))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 56), fill=PURPLE)
    d.text((32, 28), "Linx UX  ·  Cadastro de usuário de autenticação", font=F(FONT, 20), fill=WHITE, anchor="lm")
    rr(d, (80, 100, 1360, 780), 8, WHITE, (220, 220, 220), 1)
    d.text((112, 130), "Dados de autenticação", font=F(FONTB, 22), fill=PURPLE)
    d.text((112, 190), "Nome de autenticação", font=F(FONT, 15), fill=(90, 90, 90))
    rr(d, (380, 182, 1280, 226), 3, WHITE, LINE, 1)
    d.text((394, 204), "joao.silva", font=F(FONT, 17), fill=(40, 40, 40), anchor="lm")
    purple_btn(d, (112, 680, 300, 728), "Revogar MFA", ghost=True)
    purple_btn(d, (320, 680, 520, 728), "Revogar SSO", ghost=True)
    overlay = Image.new("RGBA", (W, H), (20, 16, 32, 110))
    img = img.convert("RGBA")
    img = Image.alpha_composite(img, overlay)
    d = ImageDraw.Draw(img)
    rr(d, (400, 280, 1040, 540), 8, WHITE, (200, 200, 200), 1)
    d.text((720, 320), "Informação", font=F(FONTB, 24), fill=PURPLE, anchor="mm")
    d.text((440, 370), "SSO revogado. No próximo login Microsoft", font=F(FONT, 17), fill=(50, 50, 50))
    d.text((440, 398), "o usuário vinculará a conta de novo.", font=F(FONT, 17), fill=(50, 50, 50))
    purple_btn(d, (620, 450, 820, 498), "Ok")
    img = img.convert("RGB")
    p = IMG / "tela-revogar-sso-ok.png"
    img.save(p, "PNG")
    return p


def set_run_font(run, name="Calibri", size=11, bold=False, color=SLATE):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "CBD5E1")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, "Calibri", 18 if level == 1 else 13, True, NAVY if level == 1 else PURPLE_RGB)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, "Calibri", 11, False, SLATE)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, "Calibri", 11, False, SLATE)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    set_run_font(run, "Calibri", 9, False, MUTED)


def add_picture(doc, path, width_cm=16.2, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        add_caption(doc, caption)


def add_table(doc, header, rows, col_cm):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, "Calibri", 10, True, RGBColor(255, 255, 255))
        shade(cell, "5B2E90")
        set_cell_border(cell)
        cell.width = Cm(col_cm[i])
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.rows[r + 1].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, "Calibri", 9.5, c == 0, SLATE)
            shade(cell, "F8FAFC" if r % 2 == 0 else "EEF2FF")
            set_cell_border(cell)
            cell.width = Cm(col_cm[c])
    doc.add_paragraph()


def build_docx(shots):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)

    p = doc.add_paragraph()
    set_run_font(p.add_run("LINX UX  ·  PORTAL"), "Calibri", 11, True, PURPLE_RGB)
    p = doc.add_paragraph()
    set_run_font(p.add_run("MFA e SSO no frontend"), "Calibri", 28, True, NAVY)
    p = doc.add_paragraph()
    set_run_font(p.add_run("Como a pessoa entra no Portal, o que a empresa precisa configurar e por quê."), "Calibri", 13, False, SLATE)
    p = doc.add_paragraph()
    set_run_font(p.add_run("Guia de alto nível  ·  telas do Portal  ·  sem cookbook de API"), "Calibri", 11, False, MUTED)

    add_body(doc, "Este documento tem três partes. A Parte 1 é a jornada no Portal (senha ou Microsoft + MFA). A Parte 1b descreve o processo completo de SSO e o processo de Revogar SSO. A Parte 2 é para TI da empresa: Azure, Web.config e cadastro Linx.")

    add_heading(doc, "Parte 1 — Como funciona no Portal (usuário)", 1)
    add_body(doc, "O acesso ao Linx UX no browser não é uma tela só. A pessoa identifica o login, prova quem é (senha Linx ou conta Microsoft) e, se a empresa estiver com MFA ligado, confirma um código de 6 dígitos do aplicativo de autenticação. Só então o Portal abre o Application.")
    add_body(doc, "O botão “entrar com Microsoft” prova a identidade corporativa. Ele não substitui o código TOTP do Linx. São dois fatores diferentes: a Microsoft diz quem é a pessoa; o Linx confirma o segundo fator da empresa (GPECON).")
    add_picture(doc, shots["jornada"], 16.4, "Figura 1 — Ordem das telas que a pessoa percorre no Portal.")

    add_heading(doc, "1. Identificar o usuário", 2)
    add_body(doc, "A primeira tela pede só o usuário e o botão CONTINUAR. O Portal consulta o cadastro: se aquele login tem “Utiliza SSO”, se é usuário de serviço, e qual e-mail usar como dica na Microsoft. Usuário de serviço não entra pelo Portal.")
    add_picture(doc, shots["identificar"], 16.4, "Figura 2 — Tela inicial: informe o login Linx e continue.")

    add_heading(doc, "2. Senha Linx ou Microsoft", 2)
    add_body(doc, "Depois do CONTINUAR o usuário fica travado (dá para alterar). Aparecem senha, listar ambientes, entrar — e, se o cadastro tiver Utiliza SSO e o SSO estiver ligado no Portal, o botão entrar com Microsoft.")
    add_picture(doc, shots["senha"], 16.4, "Figura 3 — Segundo passo: senha local ou “entrar com Microsoft”.")
    add_bullet(doc, "Entrar / listar ambientes — usa a senha Linx.")
    add_bullet(doc, "Entrar com Microsoft — abre a página da Microsoft (sempre pede a conta de novo). O e-mail do cadastro, quando existe, é enviado como dica de conta.")
    add_bullet(doc, "Esqueci minha senha — só no caminho de senha, se o Portal estiver com essa opção.")

    add_heading(doc, "3. Tela da Microsoft", 2)
    add_body(doc, "A pessoa autentica no diretório da empresa. O Portal não mostra a senha da Microsoft. Ao voltar, o Linx grava ou confirma o vínculo daquela conta Microsoft com o login digitado no CONTINUAR — não com um login inventado a partir do e-mail.")
    add_picture(doc, shots["microsoft"], 16.4, "Figura 4 — Autenticação na Microsoft (Entra ID). Ilustração do passo; a tela real é a do tenant da empresa.")

    add_heading(doc, "4. Se a conta Microsoft não for a vinculada", 2)
    add_body(doc, "No primeiro SSO bem-sucedido o Linx grava a conta Microsoft daquela pessoa. Nas vezes seguintes, a mesma conta entra. Se outra conta Microsoft for usada, o Portal não abre sessão e volta à tela CONTINUAR com o aviso. O administrador pode Revogar SSO no cadastro; o próximo login Microsoft grava um vínculo novo.")
    add_picture(doc, shots["erro"], 16.4, "Figura 5 — Conta Microsoft diferente da vinculada: volta ao CONTINUAR, sem entrar no produto.")

    add_heading(doc, "5. Escolher o ambiente", 2)
    add_body(doc, "O GPECON e o ambiente vêm da ficha do usuário. Se houver um ambiente padrão ou um único ambiente, o Portal segue sozinho. Se houver vários sem padrão, aparece a lista.")
    add_picture(doc, shots["ambientes"], 16.4, "Figura 6 — Lista de ambientes (quando o Portal precisa da escolha).")

    add_heading(doc, "6. Verificação em duas etapas (MFA Linx)", 2)
    add_body(doc, "Esta tela só aparece depois que o GPECON é conhecido. Não é o MFA da Microsoft. Na primeira vez a pessoa escaneia o QR no Google Authenticator ou Microsoft Authenticator e confirma o primeiro código. Nas seguintes, só informa os 6 dígitos.")
    add_picture(doc, shots["mfa1"], 16.4, "Figura 7 — Primeira vez: QR Code + código de 6 dígitos.")
    add_picture(doc, shots["mfa2"], 16.4, "Figura 8 — Acessos seguintes: só o código gerado no aplicativo.")
    add_body(doc, "Com o código aceito o Portal abre o Application. Sem o código, quando o MFA da empresa está ligado, o produto não entra. O MFA pode ser pulado se a empresa desligou MFA, se o usuário está com Utiliza MFA desmarcado, ou se a autenticação for Windows. Usuário de serviço não usa esta tela: ele não entra no Portal.")

    add_heading(doc, "7. O que o administrador vê no cadastro", 2)
    add_body(doc, "No Cadastro de usuário (Application) existem as flags Utiliza SSO e Utiliza MFA, o e-mail (dica da Microsoft) e os botões Revogar MFA e Revogar SSO, um ao lado do outro.")
    add_picture(doc, shots["cadastro"], 16.4, "Figura 9 — Flags e botões no cadastro (Application).")
    add_bullet(doc, "Utiliza SSO — libera o botão Microsoft depois do CONTINUAR.")
    add_bullet(doc, "Utiliza MFA — nulo ou ligado exige TOTP; desmarcado pula o MFA daquele usuário.")
    add_bullet(doc, "E-mail — dica na Microsoft. Não é a chave de identidade no Linx.")
    add_bullet(doc, "Revogar MFA — apaga o secret; o próximo acesso mostra o QR de novo.")
    add_bullet(doc, "Revogar SSO — apaga só o vínculo Azure. Senha e TOTP permanecem.")

    add_heading(doc, "Parte 1b — Processo completo de SSO e Revogar SSO", 1)
    add_body(doc, "O botão “entrar com Microsoft” não é um clique único. Ele amarra o login digitado no CONTINUAR à conta Azure (OID + UPN), abre a sessão do Portal e só então segue para ambiente e MFA. Revogar SSO é o caminho inverso no cadastro: apaga o vínculo sem mexer em senha nem TOTP.")
    add_picture(doc, shots["jornada_sso"], 16.4, "Figura 10 — Processo SSO completo (passos 1–5 no Portal) e Revogar SSO (passo 6 no Application).")

    add_heading(doc, "8. Processo completo de SSO", 2)
    add_body(doc, "A identidade no Linx continua sendo o Nome de autenticação digitado no CONTINUAR, não o prefixo do e-mail Microsoft. O token Azure não vai ao Service.")
    add_bullet(doc, "CONTINUAR — o Portal consulta Utiliza SSO, e-mail e usuário de serviço. Usuário de serviço é recusado. Com SSO ligado, guarda o login e o e-mail (dica da Microsoft).")
    add_bullet(doc, "Entrar com Microsoft — redirect para login.microsoftonline.com com prompt=login. Se houver e-mail no cadastro, ele vai como login_hint.")
    add_bullet(doc, "Volta ao Portal — a Microsoft devolve o code. O Portal troca o code pelo perfil (OID + UPN).")
    add_bullet(doc, "Primeiro vínculo — o Service insere AZURE_OID + AZURE_UPN em TCS_USUARIO_SSO_VINCULO.")
    add_bullet(doc, "Mesma conta — atualiza só DATA_ULTIMO_LOGIN.")
    add_bullet(doc, "Conta Microsoft diferente — recusa, não abre sessão e volta ao CONTINUAR (Figura 5).")
    add_bullet(doc, "OID já de outro usuário Linx — recusa o bind.")
    add_bullet(doc, "Sessão — só depois do vínculo aceito o Service autentica sem senha e o Portal grava o cookie.")
    add_bullet(doc, "Ambiente e MFA — iguais ao caminho de senha. SSO não dispensa o código de 6 dígitos.")

    add_heading(doc, "9. Processo de Revogar SSO", 2)
    add_body(doc, "Revoga SSO fica no cadastro de usuário (Application), sempre ao lado de Revoga MFA. O ícone da barra de ferramentas também fica sempre visível nesses cadastros. Se não houver vínculo, o clique avisa; o botão não some.")
    add_bullet(doc, "Abra o cadastro do usuário (local ou de autenticação).")
    add_bullet(doc, "Clique em Revogar SSO.")
    add_bullet(doc, "Confirme a mensagem: remove só o vínculo da conta Microsoft; o próximo SSO grava um OID/UPN novo.")
    add_bullet(doc, "O Service apaga a linha em TCS_USUARIO_SSO_VINCULO. Senha Linx, secret MFA e as flags Utiliza SSO / Utiliza MFA não mudam.")
    add_bullet(doc, "No próximo CONTINUAR → Microsoft o Linx grava um vínculo novo com a conta Azure usada naquele momento.")
    add_picture(doc, shots["revogar"], 16.4, "Figura 11 — Confirmação de Revogar SSO no cadastro.")
    add_picture(doc, shots["revogar_ok"], 16.4, "Figura 12 — SSO revogado: o próximo login Microsoft cria o vínculo de novo.")
    add_body(doc, "Use quando a pessoa trocou de conta Microsoft, quando o vínculo ficou com a conta errada, ou depois de um acesso recusado (“conta Microsoft diferente da vinculada”).")

    add_heading(doc, "10. O que fica gravado em TCS_LOG_ACESSO_AUTH", 2)
    add_body(doc, "Todo o processo de SSO e o Revogar SSO escrevem na tabela TCS_LOG_ACESSO_AUTH, canal PortalSSO. MFA TOTP não grava nesta tabela. Falhas SSOI/SSOF não contam para lockout de senha.")
    add_table(
        doc,
        ["Código", "Tipo", "Processo"],
        [
            ["SSOI-IDENT", "I", "CONTINUAR com Utiliza SSO (e-mail do login_hint)"],
            ["SSOI-START", "I", "Redirect para a Microsoft"],
            ["SSOI-AZURE", "I", "Callback com UPN depois de trocar o code"],
            ["SSOI-BIND", "I", "Login local da sessão + UPN Azure"],
            ["SSOI-FIRST", "I", "Primeiro vínculo gravado (OID + UPN)"],
            ["SSOI-LINK", "I", "Mesmo OID: atualiza último login"],
            ["(vazio) / Login SSO efetuado", "S", "AuthenticatePortalSso aceito — sessão Portal"],
            ["SSOF-OFF / CONT / START", "F", "SSO desligado, contingência ou falha ao abrir a Microsoft"],
            ["SSOF-AZURE / CODE / TOKEN", "F", "Erro Azure no callback, sem code, ou falha ao trocar o token"],
            ["SSOF-BIND / FIRST", "F", "Bind falhou (login vazio, SQL, HTTP)"],
            ["SSOF-LINK", "F", "Conta Microsoft diferente da vinculada, ou OID de outro usuário"],
            ["SSOF-EXC", "F", "Exceção MSAL no callback"],
            ["SSOI-REV", "I", "Revogar SSO: apagou o vínculo (DESCRICAO inclui by=operador)"],
            ["SSOF-REV", "F", "Revogar SSO falhou (sem vínculo, usuário inexistente, SQL)"],
        ],
        [5.2, 1.6, 9.8],
    )
    add_body(doc, "Depois de SSOI-REV o próximo SSO bem-sucedido volta a gravar SSOI-FIRST. CheckPortalSsoVinculo (consulta se há vínculo) não escreve na tabela.")

    add_heading(doc, "Parte 2 — Benefícios e o que a empresa precisa configurar", 1)
    add_heading(doc, "Benefícios", 2)
    add_bullet(doc, "Uma identidade corporativa: a pessoa entra com a conta que a empresa já gerencia no Entra ID (bloqueio, desligamento, políticas Microsoft).")
    add_bullet(doc, "Menos senha Linx no dia a dia, sem abrir mão do segundo fator da empresa (TOTP Linx no GPECON).")
    add_bullet(doc, "A cada SSO o Portal pede a conta Microsoft de novo: reduz sessão “esquecida” em browser compartilhado.")
    add_bullet(doc, "Vínculo da conta: depois do primeiro SSO, outra conta Microsoft não entra no mesmo usuário Linx.")
    add_bullet(doc, "Trilha de auditoria dos passos e erros do SSO (incluindo vínculo recusado e revogação).")
    add_bullet(doc, "Contingência: se a Microsoft estiver indisponível e o Portal permitir offline, a pessoa ainda usa usuário e senha Linx.")
    add_bullet(doc, "O e-mail do cadastro antecipa a conta na tela da Microsoft, sem misturar o login Linx com o UPN.")
    add_bullet(doc, "TI da empresa controla o app no Azure; o Linx controla quem vê o botão SSO e quem precisa de TOTP.")

    add_heading(doc, "O que a empresa precisa obter no Azure (Entra ID)", 2)
    add_body(doc, "O Portal é um aplicativo Web (confidential client). A TI da empresa cria um App registration do tipo Web — não o registro de desktop. Anote e envie ao time que configura o Portal:")
    add_table(
        doc,
        ["Dado no Azure", "Onde pegar", "Para que serve no Portal"],
        [
            ["Directory (tenant) ID", "Entra ID → Visão geral", "Qual diretório Microsoft será usado (SSO_TENANT_ID)"],
            ["Application (client) ID", "App registration → Visão geral", "Identifica o app Linx Portal (SSO_CLIENT_ID)"],
            ["Object ID do aplicativo", "App registration → Visão geral", "Referência do objeto no tenant (SSO_OBJECT_ID)"],
            ["Client secret (valor)", "Certificates & secrets → New client secret", "Fecha o retorno da Microsoft. Sem o secret o SSO Web não completa (SSO_CLIENT_SECRET)"],
            ["Redirect URI", "Authentication → Web → Redirect URIs", "Tem de ser exatamente a URL do Portal + /Account/SsoCallback"],
            ["Permissão User.Read", "API permissions → Microsoft Graph", "Ler o perfil/UPN. Consentimento do admin se a política exigir"],
            ["Supported account types", "App registration", "Em geral “somente este diretório organizacional” (tenant único)"],
        ],
        [5.2, 5.6, 5.8],
    )
    add_body(doc, "A URI de redirecionamento é o item que mais quebra implantação. Precisa ser idêntica no Azure e no Portal, inclusive http/https, host, porta e caminho. Exemplos:")
    add_bullet(doc, "https://portal.empresa.com.br/Account/SsoCallback")
    add_bullet(doc, "https://qa-ux.linx.com.br/3.11-NT/Account/SsoCallback")
    add_body(doc, "O secret tem validade. Quando expirar, o CONTINUAR + Microsoft falha na volta. Gere um secret novo no Azure e atualize só o SSO_CLIENT_SECRET no Portal — não é preciso republicar as DLLs.")

    add_heading(doc, "O que configurar no Portal (Web.config / PortalSettings)", 2)
    add_body(doc, "Estas chaves ficam na seção PortalSettings do Web.config do site Portal (a pasta onde está o Web.config). Não copie um Web.config de outro ambiente por cima: o secret e as URLs são daquele IIS.")
    add_table(
        doc,
        ["Chave", "Obrigatória", "Significado"],
        [
            ["SSO_HABILITA_AUTENTICACAO", "sim", "true liga o SSO no Portal. false esconde o fluxo Microsoft."],
            ["SSO_CLIENT_ID", "se SSO on", "Application (client) ID do App registration."],
            ["SSO_TENANT_ID", "se SSO on", "Directory (tenant) ID."],
            ["SSO_CLIENT_SECRET", "se SSO on (Web)", "Valor do client secret. Pode vir da variável SI_PDR_SSO_CLIENT_SECRET se a chave estiver vazia."],
            ["SSO_OBJECT_ID", "recomendada", "Object ID do app no Entra. Referência operacional."],
            ["SSO_REDIRECT_URI", "sim na prática", "URL pública + /Account/SsoCallback. Se vazia, o Portal monta a partir de PortalUrl."],
            ["SSO_SCOPES", "não", "Padrão User.Read. Deve bater com a permissão concedida no Azure."],
            ["SSO_PERMITE_OFFLINE", "recomendada true", "true: se a Microsoft falhar, permite senha Linx."],
            ["SSO_TIMEOUT_RESPOSTA", "não", "Segundos (padrão 120). Alinhado ao padrão OmniPOS."],
            ["PortalUrl", "sim", "URL pública deste Portal. Usada no redirect e no callback padrão."],
            ["authorizationServiceAddress", "sim", "URL do Service. O Portal consulta cadastro, vínculo e MFA nele."],
        ],
        [5.4, 3.4, 7.8],
    )

    add_heading(doc, "O que configurar no cadastro Linx (por usuário)", 2)
    add_body(doc, "Além do app Azure, cada pessoa precisa de ficha. Sem isso o botão Microsoft não aparece ou o retorno volta “sem cadastro local”.")
    add_table(
        doc,
        ["Campo", "Efeito na tela"],
        [
            ["Nome de autenticação", "É o que a pessoa digita no CONTINUAR. O SSO continua com esse login, não com o prefixo do e-mail Microsoft."],
            ["E-mail", "Dica na tela da Microsoft. Se vazio e o login já tiver @, usa o próprio login."],
            ["Utiliza SSO", "Ligado: mostra “entrar com Microsoft” depois do CONTINUAR. Desligado: só senha."],
            ["Utiliza MFA", "Nulo ou ligado: pede TOTP depois do ambiente. Desligado: pula o MFA daquele usuário."],
            ["Usuário de serviço", "Portal recusa. Esse perfil é para API, não para o frontend."],
            ["Ambiente padrão / GPECON", "Define se a lista de ambientes aparece e em qual empresa o MFA vale."],
        ],
        [5.0, 11.6],
    )

    add_heading(doc, "Checklist rápido de implantação", 2)
    add_bullet(doc, "App registration Web no tenant da empresa, com secret válido e User.Read (admin consent se preciso).")
    add_bullet(doc, "Redirect URI no Azure = URL real do Portal + /Account/SsoCallback.")
    add_bullet(doc, "PortalSettings preenchidas; reciclar o pool do Portal depois de mudar o Web.config.")
    add_bullet(doc, "Service atualizado (senão o retorno da Microsoft pode mostrar “Falha ao gravar vínculo SSO: Not Found”).")
    add_bullet(doc, "Script de banco APPLY_SSO_MFA.sql no catálogo Portal (flags, vínculo e log).")
    add_bullet(doc, "Pelo menos um usuário de teste com Utiliza SSO, e-mail corporativo e cadastro local.")
    add_bullet(doc, "Prova: CONTINUAR → Microsoft → (QR se for a primeira vez) → Application.")
    add_bullet(doc, "Prova negativa: outro usuário Microsoft no mesmo login Linx deve voltar ao CONTINUAR.")

    add_heading(doc, "O que isto não é", 2)
    add_bullet(doc, "Não é o MFA da Microsoft no lugar do TOTP Linx.")
    add_bullet(doc, "Não é um único clique que escolhe ambiente sozinho — ambiente padrão é cadastro à parte.")
    add_bullet(doc, "O cookie do Portal, sozinho, não abre o Application: depois do MFA o Portal leva a pessoa ao produto.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    set_run_font(
        p.add_run(
            "As telas reproduzem o fluxo e o visual do Portal Linx UX (roxo #5B2E90, identifier-first, botão Microsoft, MFA e cadastro). "
            "A tela da Microsoft é ilustrativa do passo; a real é a do tenant da empresa. "
            "Detalhe de APIs, se necessário: docs/login-mfa-sso-api.md."
        ),
        "Calibri",
        9,
        False,
        MUTED,
    )

    ART.mkdir(parents=True, exist_ok=True)
    out1 = ART / DOCX_NAME
    out2 = DOCS / DOCX_NAME
    doc.save(str(out1))
    shutil.copy2(out1, out2)
    print("DOCX", out1, out1.stat().st_size)
    return out1


def main():
    IMG.mkdir(parents=True, exist_ok=True)
    shots = {
        "identificar": shot_identificar(),
        "senha": shot_senha(),
        "microsoft": shot_microsoft(),
        "mfa1": shot_mfa(True),
        "mfa2": shot_mfa(False),
        "ambientes": shot_ambientes(),
        "cadastro": shot_cadastro(),
        "erro": shot_erro(),
        "jornada": shot_jornada(),
        "jornada_sso": shot_jornada_sso(),
        "revogar": shot_revogar_confirm(),
        "revogar_ok": shot_revogar_ok(),
    }
    dest = ART / "mfa-sso-docx-prints"
    dest.mkdir(parents=True, exist_ok=True)
    for p in shots.values():
        shutil.copy2(p, dest / p.name)
    build_docx(shots)


if __name__ == "__main__":
    main()
