#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera DOCX com tarefas e cenários de teste da feature MFA/SSO vs SI-PDR-CICD-AWS."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path("/workspace/docs")
ART = Path("/opt/cursor/artifacts")
DOCX_NAME = "Linx-UX-Tarefas-Cenarios-MFA-SSO-vs-SI-PDR.docx"

NAVY = RGBColor(0x0B, 0x25, 0x45)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
HEADER_BG = "0B2545"
ALT_ROW = "F1F5F9"

FEATURE = (
    "feature/implementacao-completa-ss-mfa-5-10-2026-mais-correcao-tela-senha-"
    "ficando-em-branco-antes-e-depois-n-janela-mfa-mudanca-envio-senhaget-post"
)
BASE = "SI-PDR-CICD-AWS"


def set_run(run, *, size=11, bold=False, color=SLATE):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, size=16 if level == 1 else 13, bold=True, color=NAVY)
    return p


def add_para(doc, text, *, size=11, bold=False, color=SLATE, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_text(cell, text, *, bold=False, color=SLATE, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10)
        shade_cell(table.rows[0].cells[i], HEADER_BG)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], val, size=10)
            if r_idx % 2 == 1:
                shade_cell(table.rows[r_idx + 1].cells[c_idx], ALT_ROW)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run(run, size=11, color=SLATE)
        p.paragraph_format.space_after = Pt(3)


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("Linx UX — Tarefas e cenários de teste")
    set_run(r, size=22, bold=True, color=NAVY)
    title.paragraph_format.space_after = Pt(4)

    add_para(
        doc,
        "Comparação da branch de feature MFA/SSO em relação à base SI-PDR-CICD-AWS.",
        size=11,
        color=MUTED,
        space_after=10,
    )

    add_heading(doc, "1. Escopo da comparação", level=1)
    add_table(
        doc,
        ["Item", "Valor"],
        [
            ["Branch base", BASE],
            ["Branch feature", FEATURE],
            ["Commits exclusivos da feature", "30"],
            ["Arquivos de produto (Main/)", "~45"],
            ["Observação", "O restante do diff é pacote IIS, docs e CI"],
        ],
    )
    add_para(
        doc,
        "A base SI-PDR-CICD-AWS já continha MFA. A feature completa o vínculo SSO, "
        "correções de UX (tela de senha em branco), janela TOTP configurável e o envio "
        "de senha do AuthenticateJson via POST.",
        space_after=12,
    )

    add_heading(doc, "2. Tarefas / entregas na feature", level=1)

    add_heading(doc, "2.1 SSO Azure — vínculo persistente e Revogar SSO", level=2)
    add_bullets(
        doc,
        [
            "Tabela LX_TCS.TCS_USUARIO_SSO_VINCULO (OID/UPN ↔ usuário Linx).",
            "No callback SSO: grava/confirma vínculo antes do cookie Forms (BindPortalSsoVinculo).",
            "Rejeita conta Microsoft diferente da vinculada (SSOF-LINK) e OID já ligado a outro usuário.",
            "API/UI Revoga SSO (não altera senha, MFA nem flags Utiliza).",
            "Em erro SSO: volta à tela de identificação limpa (ResetSsoLoginState), sem shell/senha residual.",
            "Com SSO disponível na etapa identificada: esconde o campo senha (fluxo SSO-only).",
        ],
    )

    add_heading(doc, "2.2 MFA — janela TOTP configurável", level=2)
    add_bullets(
        doc,
        [
            "AppSetting MfaTotpWindowSteps (default 2 ≈ ±90s; clamp 0–10).",
            "Antes: janela fixa ±2.",
        ],
    )

    add_heading(doc, "2.3 Correção “tela de senha em branco”", level=2)
    add_para(
        doc,
        "Antes/depois do MFA / expiração de senha:",
        bold=True,
        space_after=4,
    )
    add_bullets(
        doc,
        [
            "Modal de troca de senha obrigatória: em sucesso, cancelar ou falha de soft-logoff → redirect para login, sem shell Durandal vazio.",
            "Guards pagehide / detached para não ficar tela sem alvos de clique.",
        ],
    )

    add_heading(doc, "2.4 AuthenticateJson: senha GET → POST", level=2)
    add_bullets(
        doc,
        [
            "Clientes (GeranciamentoToken, WebClientHelper, LocalServiceBus, etc.) enviam userName/password/applicationId no body JSON via POST.",
            "DomainService: HasSideEffects = true para aceitar invoke com side effect (POST).",
        ],
    )

    add_heading(doc, "2.5 Cadastro de usuário — senha aleatória e UI", level=2)
    add_bullets(
        doc,
        [
            "Remoção/ocultação dos campos de senha no novo usuário; criação gera senha complexa e exige e-mail válido.",
            "Botão Enviar senha aleatória (API SendRandomUserPassword) + ícone na toolbar.",
            "Controles Revoga MFA / Revoga SSO / Enviar senha na mesma linha (layout horizontal).",
            "Botões habilitados só com flags/estado coerentes (Utiliza + vínculo/secret).",
        ],
    )

    add_heading(doc, "2.6 Lookup Ambientes (usuário em edição)", level=2)
    add_bullets(
        doc,
        [
            "LookUpTcsAmbiente2 usa o usuário editado (não só o logado), filtra por IdLinx e exclui ambientes já atribuídos.",
        ],
    )

    add_heading(doc, "2.7 Empacotamento / docs / CI", level=2)
    add_bullets(
        doc,
        [
            "Kits INSTALL_MFA_SSO / MFA-SSO-UPDATE, scripts SQL, guias usuário/API.",
            "Workflow AWS IIS aponta para essa branch; smoke SSO não derruba o deploy.",
        ],
    )

    add_heading(doc, "3. Cenários de teste sugeridos", level=1)

    add_heading(doc, "A. Login senha + MFA", level=2)
    add_table(
        doc,
        ["#", "Cenário", "Esperado"],
        [
            ["A1", "Usuário/senha, empresa com MFA, 1º acesso", "QR → confirma 6 dígitos → Application"],
            ["A2", "Já enrollado, código válido na janela", "Ticket MFA → Application"],
            ["A3", "Código fora da janela / inválido", "Erro; após 5 falhas, bloqueio ~15 min"],
            ["A4", "MfaTotpWindowSteps=0 vs 2", "Aceita só step atual vs ±2"],
            ["A5", "Utiliza MFA desligado ou empresa MFA off", "Entra sem desafio"],
            ["A6", "Usuário de serviço / API", "Sem MFA no Portal"],
        ],
    )

    add_heading(doc, "B. SSO + vínculo", level=2)
    add_table(
        doc,
        ["#", "Cenário", "Esperado"],
        [
            ["B1", "1º SSO (Utiliza SSO)", "Vínculo SSOI-FIRST; depois ambiente + MFA"],
            ["B2", "Mesma conta Microsoft de novo", "SSOI-LINK; atualiza último login"],
            ["B3", "Outra conta Microsoft no mesmo login Linx", "Recusa SSOF-LINK; volta ao CONTINUAR; sem cookie"],
            ["B4", "OID já ligado a outro usuário", "Recusa bind"],
            ["B5", "Erro Azure / abort", "Mensagem + tela identificação limpa; sem senha residual"],
            ["B6", "“alterar usuário”", "Reset estado SSO/MFA session"],
            ["B7", "Utiliza SSO ligado", "Após CONTINUAR, botão Microsoft e sem campo senha"],
        ],
    )

    add_heading(doc, "C. Revoga MFA / SSO (Application)", level=2)
    add_table(
        doc,
        ["#", "Cenário", "Esperado"],
        [
            ["C1", "Revoga MFA com secret", "Limpa secret; próximo login pede QR de novo"],
            ["C2", "Revoga SSO com vínculo", "Apaga linha vínculo; senha/MFA/flags intactos; próximo SSO revincula"],
            ["C3", "Sem vínculo / sem Utiliza", "Botão visível; enable/aviso conforme regra"],
            ["C4", "Cadastro Local, Autenticação e Manutenção", "Mesma linha: MFA + SSO (+ senha aleatória)"],
        ],
    )

    add_heading(doc, "D. Senha em branco / expiração", level=2)
    add_table(
        doc,
        ["#", "Cenário", "Esperado"],
        [
            ["D1", "Senha expirada → modal → altera com sucesso", "Redirect login; não fica tela branca"],
            ["D2", "Cancelar / fechar no fluxo obrigatório", "Volta ao login"],
            ["D3", "Soft-logoff falha", "Redirect login (não shell vazio)"],
            ["D4", "Fluxo após MFA (se cair em change-password)", "Mesmo comportamento de saída limpa"],
        ],
    )

    add_heading(doc, "E. AuthenticateJson POST", level=2)
    add_table(
        doc,
        ["#", "Cenário", "Esperado"],
        [
            ["E1", "Desktop/LocalServiceBus login válido", "POST body autentica; token ok"],
            ["E2", "Senha com caracteres especiais (&, +, espaço)", "Funciona (antes quebrava no GET)"],
            ["E3", "GET antigo com querystring", "Deve falhar ou não ser o caminho usado"],
            ["E4", "TcsAutorizacao e Autorizacao DomainService", "Ambos via POST"],
        ],
    )

    add_heading(doc, "F. Novo usuário / senha aleatória", level=2)
    add_table(
        doc,
        ["#", "Cenário", "Esperado"],
        [
            ["F1", "Novo usuário sem campos de senha", "UI sem digitar senha; e-mail obrigatório"],
            ["F2", "Criar com e-mail inválido", "Erro; não cria"],
            ["F3", "Criar com e-mail válido", "Senha complexa gerada + e-mail"],
            ["F4", "Enviar senha aleatória", "Nova senha, e-mail, expiração no próximo acesso"],
            ["F5", "Sem e-mail", "Botão desabilitado / mensagem"],
        ],
    )

    add_heading(doc, "G. Ambientes no cadastro", level=2)
    add_table(
        doc,
        ["#", "Cenário", "Esperado"],
        [
            ["G1", "Editar usuário e abrir lookup Ambientes", "Lista do usuário editado, não do admin logado"],
            ["G2", "Ambientes já na aba", "Não reaparecem no lookup"],
            ["G3", "Filtro IdLinx", "Só ambientes da empresa correta"],
        ],
    )

    add_heading(doc, "H. Regressão rápida", level=2)
    add_bullets(
        doc,
        [
            "Login só senha sem SSO.",
            "MFA skip documentado (Windows auth / flags).",
            "Deploy AWS: Portal/Service/Application sobem; smoke SSO não falha o pipeline.",
        ],
    )

    add_heading(doc, "4. Ordem prática de QA", level=1)
    add_para(
        doc,
        "Schema SQL → F (cadastro/senha) → A (senha+MFA) → B (SSO) → C (revoga) → "
        "D (tela branca) → E (POST desktop) → G (ambientes).",
        space_after=12,
    )

    add_heading(doc, "5. Fontes", level=1)
    add_bullets(
        doc,
        [
            f"Diff git: origin/{BASE}...origin/{FEATURE}",
            "docs/login-mfa-sso-usuario.md",
            "docs/estorias-construcao-mfa.md",
        ],
    )

    ART.mkdir(parents=True, exist_ok=True)
    DOCS.mkdir(parents=True, exist_ok=True)
    out_docs = DOCS / DOCX_NAME
    out_art = ART / DOCX_NAME
    doc.save(out_docs)
    shutil.copy2(out_docs, out_art)
    print(f"Wrote {out_docs}")
    print(f"Copied {out_art}")


if __name__ == "__main__":
    build()
