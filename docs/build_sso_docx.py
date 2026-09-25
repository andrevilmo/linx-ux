#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""DOCX de alto nível só de SSO (usuário + requisitos da empresa)."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

IMG = Path("/tmp/sso-docx/img")
ART = Path("/opt/cursor/artifacts")
DOCS = Path("/workspace/docs")
DOCX_NAME = "Linx-UX-SSO-Portal.docx"

FONT = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONTB = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
FONTI = "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf"

PURPLE = (91, 46, 144)
WHITE = (255, 255, 255)
GRAY = (117, 117, 117)
LINE = (207, 207, 207)
NAVY_T = (11, 37, 69)
SLATE_T = (51, 65, 85)
BG_L = (32, 18, 58)
ERR = (198, 40, 40)

NAVY = RGBColor(0x0B, 0x25, 0x45)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
PURPLE_RGB = RGBColor(0x5B, 0x2E, 0x90)


def F(path, size):
    return ImageFont.truetype(path, size)


def rr(d, box, r, fill, outline=None, w=2):
    d.rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=w)


def portal_frame():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), WHITE)
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 300, H), fill=BG_L)
    for i in range(20):
        d.rectangle((0, i * 44, 300, i * 44 + 22), fill=(42, 24, 72))
    d.text((28, 36), "LINX", font=F(FONTB, 34), fill=WHITE)
    d.text((28, 92), "SOFTWARE QUE", font=F(FONT, 16), fill=(230, 220, 245))
    d.text((28, 116), "MOVE O VAREJO", font=F(FONT, 16), fill=(230, 220, 245))
    d.text((16, H - 36), "© Linx - Todos os direitos reservados.", font=F(FONT, 12), fill=(200, 190, 210))
    d.text((360, 70), "linx", font=F(FONTB, 56), fill=(30, 30, 30))
    tw = d.textlength("linx", font=F(FONTB, 56))
    d.text((360 + tw, 70), "ux", font=F(FONTB, 56), fill=PURPLE)
    d.text((W - 360, H - 28), "© Linx - Todos os direitos reservados.", font=F(FONT, 13), fill=GRAY)
    return img, d


def underline_field(d, x, y, w, label, value, readonly=False):
    d.text((x, y), label.upper(), font=F(FONT, 15), fill=GRAY)
    d.text((x, y + 32), value, font=F(FONT, 22), fill=(40, 40, 40) if value else (180, 180, 180))
    d.line((x, y + 68, x + w, y + 68), fill=PURPLE if readonly else LINE, width=2)


def purple_btn(d, box, text, ghost=False):
    if ghost:
        rr(d, box, 2, (238, 238, 238), PURPLE, 1)
        d.text(((box[0] + box[2]) / 2, (box[1] + box[3]) / 2), text.upper(), font=F(FONT, 16), fill=PURPLE, anchor="mm")
    else:
        rr(d, box, 2, PURPLE)
        d.text(((box[0] + box[2]) / 2, (box[1] + box[3]) / 2), text.upper(), font=F(FONT, 16), fill=WHITE, anchor="mm")


def shot_identificar():
    img, d = portal_frame()
    underline_field(d, 360, 240, 680, "usuário", "")
    purple_btn(d, (360, 360, 1040, 410), "continuar")
    p = IMG / "tela-identificar.png"
    img.save(p, "PNG")
    return p


def shot_senha():
    img, d = portal_frame()
    underline_field(d, 360, 210, 680, "usuário", "joao.silva", True)
    d.text((360, 290), "alterar usuário", font=F(FONT, 16), fill=PURPLE)
    underline_field(d, 360, 330, 680, "senha", "••••••••")
    purple_btn(d, (360, 460, 740, 508), "listar ambientes", ghost=True)
    purple_btn(d, (760, 460, 1040, 508), "entrar")
    d.text((700, 540), "ou", font=F(FONT, 16), fill=(160, 160, 160), anchor="mm")
    purple_btn(d, (360, 568, 1040, 620), "entrar com Microsoft")
    p = IMG / "tela-senha-sso.png"
    img.save(p, "PNG")
    return p


def shot_microsoft():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), (242, 242, 242))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 52), fill=WHITE)
    d.line((0, 52, W, 52), fill=(225, 225, 225), width=1)
    d.rectangle((24, 14, 34, 24), fill=(242, 80, 34))
    d.rectangle((36, 14, 46, 24), fill=(127, 186, 0))
    d.rectangle((24, 26, 34, 36), fill=(0, 164, 239))
    d.rectangle((36, 26, 46, 36), fill=(255, 185, 0))
    d.text((58, 26), "login.microsoftonline.com", font=F(FONT, 16), fill=(97, 97, 97), anchor="lm")
    rr(d, (500, 160, 940, 620), 2, WHITE, (220, 220, 220), 1)
    d.rectangle((532, 196, 542, 206), fill=(242, 80, 34))
    d.rectangle((544, 196, 554, 206), fill=(127, 186, 0))
    d.rectangle((532, 208, 542, 218), fill=(0, 164, 239))
    d.rectangle((544, 208, 554, 218), fill=(255, 185, 0))
    d.text((532, 250), "Entrar", font=F(FONTB, 32), fill=(27, 27, 27))
    d.text((532, 300), "joao.silva@empresa.com", font=F(FONT, 18), fill=(97, 97, 97))
    d.text((532, 380), "Senha", font=F(FONT, 16), fill=(141, 141, 141))
    d.line((532, 420, 908, 420), fill=(138, 136, 134), width=1)
    rr(d, (760, 470, 908, 514), 2, (0, 103, 184))
    d.text((834, 492), "Avançar", font=F(FONT, 16), fill=WHITE, anchor="mm")
    p = IMG / "tela-microsoft.png"
    img.save(p, "PNG")
    return p


def shot_ambientes():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), (26, 16, 48))
    d = ImageDraw.Draw(img)
    d.text((48, 36), "LINX", font=F(FONTB, 28), fill=WHITE)
    d.text((48, 120), "Ambientes disponíveis:", font=F(FONT, 26), fill=WHITE)
    d.rectangle((48, 180, 1390, 248), fill=PURPLE)
    for i, h in enumerate(["Ambiente", "Empresa", "Grupo econômico", ""]):
        d.text((70 + i * 330, 214), h, font=F(FONTB, 16), fill=WHITE, anchor="lm")
    rows = [("Produção", "ACME Varejo Ltda", "ACME"), ("Homologação", "ACME Varejo Ltda", "ACME")]
    for r, row in enumerate(rows):
        y = 248 + r * 64
        d.rectangle((48, y, 1390, y + 64), fill=WHITE if r == 0 else (248, 248, 252))
        for i, val in enumerate(row):
            d.text((70 + i * 330, y + 32), val, font=F(FONT, 18), fill=(60, 60, 60), anchor="lm")
        d.text((1220, y + 32), "entrar", font=F(FONT, 18), fill=PURPLE, anchor="lm")
    p = IMG / "tela-ambientes.png"
    img.save(p, "PNG")
    return p


def shot_cadastro():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), (245, 246, 248))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 56), fill=PURPLE)
    d.text((32, 28), "Linx UX  ·  Cadastro de usuário", font=F(FONT, 20), fill=WHITE, anchor="lm")
    rr(d, (80, 100, 1360, 780), 8, WHITE, (220, 220, 220), 1)
    d.text((112, 130), "Requisitos no cadastro (SSO)", font=F(FONTB, 22), fill=PURPLE)
    fields = [
        ("Nome de autenticação", "joao.silva"),
        ("E-mail", "joao.silva@empresa.com"),
    ]
    y = 190
    for lab, val in fields:
        d.text((112, y), lab, font=F(FONT, 15), fill=(90, 90, 90))
        rr(d, (380, y - 8, 1280, y + 36), 3, WHITE, LINE, 1)
        d.text((394, y + 14), val, font=F(FONT, 17), fill=(40, 40, 40), anchor="lm")
        y += 64
    for label, on in (("Utiliza SSO", True), ("Usuário de serviço", False)):
        rr(d, (112, y, 136, y + 24), 3, PURPLE if on else WHITE, PURPLE, 2)
        if on:
            d.line((118, y + 13, 124, y + 19), fill=WHITE, width=2)
            d.line((124, y + 19, 132, y + 7), fill=WHITE, width=2)
        d.text((148, y + 12), label, font=F(FONT, 17), fill=(50, 50, 50), anchor="lm")
        y += 40
    d.text((112, y + 8), "Utiliza SSO ligado + e-mail corporativo. Usuário de serviço não entra no Portal.", font=F(FONTI, 15), fill=GRAY)
    d.text((112, y + 36), "Revoga SSO fica sempre ao lado de Revoga MFA.", font=F(FONTI, 15), fill=GRAY)
    purple_btn(d, (112, 680, 300, 728), "Revoga MFA", ghost=True)
    purple_btn(d, (320, 680, 520, 728), "Revoga SSO", ghost=True)
    p = IMG / "tela-cadastro-sso.png"
    img.save(p, "PNG")
    return p


def shot_erro():
    img, d = portal_frame()
    d.text((360, 200), "Conta Microsoft diferente da vinculada a este usuário.", font=F(FONT, 18), fill=ERR)
    d.text((360, 228), "Use a conta Azure já associada ou peça para revogar o SSO no cadastro.", font=F(FONT, 16), fill=ERR)
    underline_field(d, 360, 300, 680, "usuário", "joao.silva")
    purple_btn(d, (360, 420, 1040, 470), "continuar")
    p = IMG / "tela-vinculo-erro.png"
    img.save(p, "PNG")
    return p


def shot_jornada():
    W, H = 1600, 420
    img = Image.new("RGB", (W, H), (244, 247, 251))
    d = ImageDraw.Draw(img)
    d.text((60, 24), "SSO no Portal — o que a pessoa percorre", font=F(FONTB, 28), fill=NAVY_T)
    steps = [
        ("1", "Usuário", "CONTINUAR", (13, 115, 119)),
        ("2", "Microsoft", "entrar com Microsoft", PURPLE),
        ("3", "Vínculo", "OID + UPN", (27, 79, 138)),
        ("4", "Ambiente", "se houver vários", (224, 122, 61)),
        ("5", "Application", "produto aberto", (46, 125, 79)),
    ]
    for i, (n, a, b, c) in enumerate(steps):
        x = 50 + i * 310
        rr(d, (x, 110, x + 270, 360), 18, WHITE, c, 4)
        d.ellipse((x + 20, 132, x + 78, 190), fill=c)
        d.text((x + 49, 161), n, font=F(FONTB, 26), fill=WHITE, anchor="mm")
        d.text((x + 24, 220), a, font=F(FONTB, 24), fill=NAVY_T)
        d.text((x + 24, 268), b, font=F(FONT, 20), fill=SLATE_T)
        if i < 4:
            d.polygon([(x + 278, 235), (x + 302, 225), (x + 302, 245)], fill=c)
    p = IMG / "jornada-sso-usuario.png"
    img.save(p, "PNG")
    return p


def shot_revogar_confirm():
    W, H = 1440, 860
    img = Image.new("RGB", (W, H), (245, 246, 248))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 56), fill=PURPLE)
    d.text((32, 28), "Linx UX  ·  Cadastro de usuário", font=F(FONT, 20), fill=WHITE, anchor="lm")
    rr(d, (80, 100, 1360, 780), 8, WHITE, (220, 220, 220), 1)
    purple_btn(d, (112, 680, 300, 728), "Revoga MFA", ghost=True)
    purple_btn(d, (320, 680, 520, 728), "Revoga SSO", ghost=True)
    overlay = Image.new("RGBA", (W, H), (20, 16, 32, 110))
    img = img.convert("RGBA")
    img = Image.alpha_composite(img, overlay)
    d = ImageDraw.Draw(img)
    rr(d, (380, 250, 1060, 560), 8, WHITE, (200, 200, 200), 1)
    d.text((720, 292), "Revoga SSO", font=F(FONTB, 24), fill=PURPLE, anchor="mm")
    d.text((420, 340), "Revogar o SSO deste usuário? Remove só o vínculo", font=F(FONT, 17), fill=(50, 50, 50))
    d.text((420, 368), "da conta Microsoft. O próximo login SSO gravará", font=F(FONT, 17), fill=(50, 50, 50))
    d.text((420, 396), "um novo OID/UPN.", font=F(FONT, 17), fill=(50, 50, 50))
    purple_btn(d, (520, 470, 700, 518), "Não", ghost=True)
    purple_btn(d, (740, 470, 920, 518), "Sim")
    img = img.convert("RGB")
    p = IMG / "tela-revogar-sso.png"
    img.save(p, "PNG")
    return p


def set_run_font(run, name="Calibri", size=11, bold=False, color=SLATE):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "CBD5E1")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, "Calibri", 18 if level == 1 else 13, True, NAVY if level == 1 else PURPLE_RGB)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, "Calibri", 11, False, SLATE)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, "Calibri", 11, False, SLATE)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    set_run_font(run, "Calibri", 9, False, MUTED)


def add_picture(doc, path, width_cm=16.2, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        add_caption(doc, caption)


def add_table(doc, header, rows, col_cm):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, "Calibri", 10, True, RGBColor(255, 255, 255))
        shade(cell, "5B2E90")
        set_cell_border(cell)
        cell.width = Cm(col_cm[i])
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.rows[r + 1].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, "Calibri", 9.5, c == 0, SLATE)
            shade(cell, "F8FAFC" if r % 2 == 0 else "EEF2FF")
            set_cell_border(cell)
            cell.width = Cm(col_cm[c])
    doc.add_paragraph()


def build_docx(shots):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)

    p = doc.add_paragraph()
    set_run_font(p.add_run("LINX UX  ·  PORTAL"), "Calibri", 11, True, PURPLE_RGB)
    p = doc.add_paragraph()
    set_run_font(p.add_run("SSO — entrar com Microsoft"), "Calibri", 28, True, NAVY)
    p = doc.add_paragraph()
    set_run_font(p.add_run("Guia de alto nível para quem usa o Portal e para quem libera o SSO na empresa."), "Calibri", 13, False, SLATE)
    p = doc.add_paragraph()
    set_run_font(p.add_run("Somente SSO  ·  telas do Portal  ·  requisitos  ·  sem cookbook de API e sem MFA"), "Calibri", 11, False, MUTED)

    add_body(doc, "Este documento trata só do login Microsoft (SSO). Não descreve o MFA Linx. Se a empresa tiver MFA ligado, o código de 6 dígitos continua existindo depois do SSO — o botão Microsoft não o substitui.")

    add_heading(doc, "Parte 1 — Como a pessoa entra", 1)
    add_picture(doc, shots["jornada"], 16.4, "Figura 1 — Ordem das telas do SSO no Portal.")

    add_heading(doc, "1. CONTINUAR com o login Linx", 2)
    add_body(doc, "A primeira tela pede só o usuário e o botão CONTINUAR. A identidade no Linx é esse Nome de autenticação — não o prefixo do e-mail Microsoft. O Portal consulta se o cadastro tem Utiliza SSO, se é usuário de serviço e qual e-mail usar como dica. Usuário de serviço não entra pelo Portal.")
    add_picture(doc, shots["identificar"], 16.4, "Figura 2 — Informe o login Linx e continue.")

    add_heading(doc, "2. Entrar com Microsoft", 2)
    add_body(doc, "Se Utiliza SSO estiver ligado e o Portal tiver SSO habilitado, aparece o botão entrar com Microsoft. A Microsoft sempre pede a conta de novo. O e-mail do cadastro, quando existe, vai como dica de conta.")
    add_picture(doc, shots["senha"], 16.4, "Figura 3 — Depois do CONTINUAR: senha Linx ou entrar com Microsoft.")
    add_picture(doc, shots["microsoft"], 16.4, "Figura 4 — Tela da Microsoft (Entra ID). Ilustração do passo; a tela real é a do tenant da empresa.")

    add_heading(doc, "3. Vínculo da conta Microsoft", 2)
    add_body(doc, "Ao voltar, o Linx amarra a conta Microsoft (OID + UPN) ao login digitado no CONTINUAR. O token Azure não vai ao Service.")
    add_bullet(doc, "Primeira vez — grava a conta Microsoft naquele usuário.")
    add_bullet(doc, "Mesma conta — entra e atualiza a data do último login.")
    add_bullet(doc, "Outra conta Microsoft no mesmo login Linx — recusa, não abre sessão e volta ao CONTINUAR.")
    add_bullet(doc, "A mesma conta Microsoft já ligada a outro usuário Linx — recusa.")
    add_picture(doc, shots["erro"], 16.4, "Figura 5 — Conta Microsoft diferente da vinculada: volta ao CONTINUAR.")

    add_heading(doc, "4. Ambiente", 2)
    add_body(doc, "O GPECON e o ambiente vêm da ficha. Um ambiente padrão ou um único ambiente seguem sozinhos. Vários sem padrão pedem escolha.")
    add_picture(doc, shots["ambientes"], 16.4, "Figura 6 — Lista de ambientes, quando o Portal precisa da escolha.")

    add_heading(doc, "5. Revoga SSO", 2)
    add_body(doc, "No cadastro de usuário (Application), Revoga SSO fica sempre ao lado de Revoga MFA. O ícone da barra também fica sempre visível nesses cadastros. Remove só o vínculo. Senha Linx e a flag Utiliza SSO não mudam. No próximo CONTINUAR → Microsoft o Linx grava um vínculo novo.")
    add_picture(doc, shots["cadastro"], 16.4, "Figura 7 — Requisitos no cadastro e os dois botões sempre visíveis.")
    add_picture(doc, shots["revogar"], 16.4, "Figura 8 — Confirmação de Revoga SSO.")
    add_body(doc, "Use quando a pessoa trocou de conta Microsoft, quando o vínculo ficou errado, ou depois do aviso da Figura 5.")

    add_heading(doc, "Parte 2 — Requisitos", 1)
    add_body(doc, "O SSO só aparece e só completa se a empresa, o Portal, o cadastro e a pessoa atenderem ao que segue.")

    add_heading(doc, "Requisitos da empresa (Microsoft Entra ID)", 2)
    add_body(doc, "O Portal é um aplicativo Web (confidential client). A TI cria um App registration do tipo Web — não o registro de desktop.")
    add_table(
        doc,
        ["Requisito", "Onde", "Para que serve"],
        [
            ["Tenant Entra ID da empresa", "Entra ID → Visão geral", "Diretório onde as pessoas autenticam"],
            ["App registration tipo Web", "App registrations → New", "O Portal não usa o registro de desktop"],
            ["Directory (tenant) ID", "Visão geral do diretório", "SSO_TENANT_ID"],
            ["Application (client) ID", "Visão geral do app", "SSO_CLIENT_ID"],
            ["Object ID do aplicativo", "Visão geral do app", "SSO_OBJECT_ID (referência)"],
            ["Client secret válido", "Certificates & secrets", "SSO_CLIENT_SECRET. Sem o secret o retorno não fecha"],
            ["Redirect URI idêntica", "Authentication → Web", "URL pública do Portal + /Account/SsoCallback"],
            ["Permissão User.Read", "API permissions → Graph", "Ler perfil/UPN. Admin consent se a política exigir"],
            ["Só este diretório (em geral)", "Supported account types", "Evita conta pessoal / outro tenant"],
            ["Pessoas ativas no Entra", "Usuários do tenant", "Quem não entra na Microsoft não entra no SSO Linx"],
        ],
        [5.2, 5.4, 6.0],
    )
    add_body(doc, "A URI de redirecionamento é o item que mais quebra implantação. Precisa ser idêntica no Azure e no Portal (http/https, host, porta e caminho).")
    add_bullet(doc, "https://portal.empresa.com.br/Account/SsoCallback")
    add_bullet(doc, "https://qa-ux.linx.com.br/3.11-NT/Account/SsoCallback")
    add_body(doc, "O secret expira. Gere outro no Azure e atualize só SSO_CLIENT_SECRET — não é preciso republicar DLL.")

    add_heading(doc, "Requisitos do Portal (Web.config / PortalSettings)", 2)
    add_body(doc, "Chaves na seção PortalSettings do Web.config do site Portal. Não copie um Web.config de outro ambiente por cima. Recicle o pool do Portal depois de alterar.")
    add_table(
        doc,
        ["Chave", "Obrigatória", "Significado"],
        [
            ["SSO_HABILITA_AUTENTICACAO", "sim", "true liga o SSO. false esconde o Microsoft."],
            ["SSO_CLIENT_ID", "se SSO on", "Application (client) ID."],
            ["SSO_TENANT_ID", "se SSO on", "Directory (tenant) ID."],
            ["SSO_CLIENT_SECRET", "se SSO on", "Valor do secret (ou SI_PDR_SSO_CLIENT_SECRET)."],
            ["SSO_OBJECT_ID", "recomendada", "Object ID do app no Entra."],
            ["SSO_REDIRECT_URI", "sim na prática", "URL pública + /Account/SsoCallback."],
            ["SSO_SCOPES", "não", "Padrão User.Read."],
            ["SSO_PERMITE_OFFLINE", "recomendada true", "Se a Microsoft falhar, permite senha Linx."],
            ["SSO_TIMEOUT_RESPOSTA", "não", "Segundos (padrão 120)."],
            ["PortalUrl", "sim", "URL pública deste Portal."],
            ["authorizationServiceAddress", "sim", "URL do Service (cadastro e vínculo)."],
        ],
        [5.4, 3.4, 7.8],
    )

    add_heading(doc, "Requisitos do cadastro Linx (por pessoa)", 2)
    add_table(
        doc,
        ["Campo", "Efeito"],
        [
            ["Nome de autenticação", "O que a pessoa digita no CONTINUAR. O SSO continua com esse login."],
            ["E-mail", "Dica na Microsoft. Se vazio e o login já tiver @, usa o próprio login."],
            ["Utiliza SSO", "Ligado: mostra entrar com Microsoft. Desligado: só senha."],
            ["Usuário de serviço", "Portal recusa. Esse perfil é para API."],
            ["Ambiente padrão / GPECON", "Define se a lista de ambientes aparece."],
        ],
        [5.0, 11.6],
    )
    add_body(doc, "Sem ficha local o retorno da Microsoft volta “sem cadastro local”. Sem Utiliza SSO o botão Microsoft não aparece.")

    add_heading(doc, "Requisitos do Service e do banco", 2)
    add_bullet(doc, "Service atualizado com vínculo SSO. Sem isso o callback pode mostrar “Falha ao gravar vínculo SSO: Not Found”.")
    add_bullet(doc, "Script APPLY_SSO_MFA.sql no catálogo Portal (tabela TCS_USUARIO_SSO_VINCULO e log TCS_LOG_ACESSO_AUTH).")
    add_bullet(doc, "Não rode o script no catálogo da Application.")

    add_heading(doc, "Requisitos de quem usa o Portal", 2)
    add_bullet(doc, "Browser (não é o caminho do usuário de serviço).")
    add_bullet(doc, "Saber o login Linx para o CONTINUAR.")
    add_bullet(doc, "Conta Microsoft corporativa ativa no tenant da empresa.")

    add_heading(doc, "Parte 3 — Benefícios, o que não é e checklist", 1)
    add_heading(doc, "Benefícios", 2)
    add_bullet(doc, "A pessoa entra com a conta que a empresa já gerencia no Entra ID (bloqueio, desligamento, políticas).")
    add_bullet(doc, "Menos senha Linx no dia a dia.")
    add_bullet(doc, "A cada SSO o Portal pede a conta Microsoft de novo (browser compartilhado).")
    add_bullet(doc, "Depois do primeiro SSO, outra conta Microsoft não entra no mesmo usuário Linx.")
    add_bullet(doc, "Revoga SSO troca o vínculo sem apagar senha.")
    add_bullet(doc, "Se a Microsoft estiver fora e SSO_PERMITE_OFFLINE estiver ligado, a senha Linx continua valendo.")

    add_heading(doc, "O que isto não é", 2)
    add_bullet(doc, "Não é o MFA da Microsoft no lugar de outro fator Linx.")
    add_bullet(doc, "Não escolhe ambiente sozinho — ambiente padrão é cadastro à parte.")
    add_bullet(doc, "O token Azure não vai ao Service. O Linx só usa OID + UPN para amarrar a conta.")
    add_bullet(doc, "Usuário de serviço não entra pelo Portal.")

    add_heading(doc, "Checklist", 2)
    add_bullet(doc, "App registration Web no tenant, secret válido, User.Read (admin consent se preciso).")
    add_bullet(doc, "Redirect URI no Azure = URL real do Portal + /Account/SsoCallback.")
    add_bullet(doc, "PortalSettings preenchidas; reciclar o pool do Portal.")
    add_bullet(doc, "Service com vínculo SSO; SQL APPLY_SSO_MFA.sql no catálogo Portal.")
    add_bullet(doc, "Usuário de teste: cadastro local, Utiliza SSO, e-mail corporativo, não é usuário de serviço.")
    add_bullet(doc, "Prova: CONTINUAR → Microsoft → Application (ou ambiente).")
    add_bullet(doc, "Prova negativa: outra conta Microsoft no mesmo login Linx deve voltar ao CONTINUAR.")
    add_bullet(doc, "Prova Revoga SSO: o botão aparece ao lado de Revoga MFA; o próximo Microsoft grava vínculo novo.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    set_run_font(
        p.add_run(
            "As telas reproduzem o visual do Portal Linx UX. A tela da Microsoft é ilustrativa; a real é a do tenant da empresa. "
            "Guia conjunto MFA+SSO: docs/Linx-UX-MFA-SSO-Portal.docx. Texto: docs/login-sso-usuario.md."
        ),
        "Calibri",
        9,
        False,
        MUTED,
    )

    ART.mkdir(parents=True, exist_ok=True)
    out1 = ART / DOCX_NAME
    out2 = DOCS / DOCX_NAME
    doc.save(str(out1))
    shutil.copy2(out1, out2)
    print("DOCX", out1, out1.stat().st_size)
    return out1


def main():
    IMG.mkdir(parents=True, exist_ok=True)
    shots = {
        "identificar": shot_identificar(),
        "senha": shot_senha(),
        "microsoft": shot_microsoft(),
        "ambientes": shot_ambientes(),
        "cadastro": shot_cadastro(),
        "erro": shot_erro(),
        "jornada": shot_jornada(),
        "revogar": shot_revogar_confirm(),
    }
    dest = ART / "sso-docx-prints"
    dest.mkdir(parents=True, exist_ok=True)
    for p in shots.values():
        shutil.copy2(p, dest / p.name)
    build_docx(shots)


if __name__ == "__main__":
    main()
