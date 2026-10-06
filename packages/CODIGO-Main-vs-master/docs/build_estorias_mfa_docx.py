#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""DOCX da estória de construção do MFA (a partir do markdown)."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path("/workspace/docs")
ART = Path("/opt/cursor/artifacts")
SRC = DOCS / "estorias-construcao-mfa.md"
DOCX_NAME = "Linx-UX-Estoria-Construcao-MFA.docx"

NAVY = RGBColor(0x0B, 0x25, 0x45)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
PURPLE = RGBColor(0x5B, 0x2E, 0x90)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, name, size, bold, color):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, hex_color):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc.append(shd)


def set_cell_border(cell):
    tc = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9E3")
        tcBorders.append(el)
    tc.append(tcBorders)


def add_runs(paragraph, text, size=11, color=SLATE, bold=False):
    parts = text.split("`")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        is_code = i % 2 == 1
        set_run_font(run, "Consolas" if is_code else "Calibri", 9.5 if is_code else size, bold and not is_code, MUTED if is_code else color)


def parse_table(lines):
    rows = []
    for line in lines:
        cols = [c.strip() for c in line.strip().strip("|").split("|")]
        if cols and set(cols[0]) <= set("-: "):
            continue
        rows.append(cols)
    return rows[0], rows[1:] if len(rows) > 1 else []


def build():
    md = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)

    i = 0
    in_code = False
    in_mermaid = False
    code_buf = []
    table_buf = []
    first_h1 = True

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        header, rows = parse_table(table_buf)
        table_buf = []
        if not header:
            return
        t = doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True
        for c, h in enumerate(header):
            cell = t.rows[0].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            set_run_font(run, "Calibri", 10, True, WHITE)
            shade(cell, "5B2E90")
            set_cell_border(cell)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                if c >= len(header):
                    break
                cell = t.rows[r + 1].cells[c]
                cell.text = ""
                run = cell.paragraphs[0].add_run(val.replace("`", ""))
                set_run_font(run, "Calibri", 9, c == 0, SLATE)
                shade(cell, "F8FAFC" if r % 2 == 0 else "EEF2FF")
                set_cell_border(cell)
        doc.add_paragraph()

    def flush_code():
        nonlocal code_buf
        if not code_buf:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run("\n".join(code_buf))
        set_run_font(run, "Consolas", 8.5, False, SLATE)
        code_buf = []

    while i < len(md):
        line = md[i]
        if line.startswith("```"):
            flush_table()
            fence = line.strip("`").strip()
            if in_code or in_mermaid:
                if in_mermaid:
                    p = doc.add_paragraph()
                    add_runs(p, "Diagrama do fluxo: ver o markdown (bloco mermaid) ou a figura do guia de APIs.", 11, MUTED)
                else:
                    flush_code()
                in_code = False
                in_mermaid = False
            else:
                in_mermaid = fence.startswith("mermaid")
                in_code = not in_mermaid
            i += 1
            continue
        if in_mermaid:
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        flush_table()
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            if first_h1:
                set_run_font(p.add_run("LINX UX  ·  DOCUMENTAÇÃO DE PRODUTO"), "Calibri", 11, True, PURPLE)
                p = doc.add_paragraph()
                first_h1 = False
            set_run_font(p.add_run(line[2:].strip()), "Calibri", 26, True, NAVY)
        elif line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            set_run_font(p.add_run(line[3:].strip()), "Calibri", 16, True, PURPLE)
        elif line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            set_run_font(p.add_run(line[4:].strip()), "Calibri", 13, True, NAVY)
        elif line.startswith("---"):
            pass
        elif line.startswith("- [x] ") or line.startswith("- [ ] "):
            done = line.startswith("- [x]")
            p = doc.add_paragraph()
            mark = "☑ " if done else "☐ "
            add_runs(p, mark + line[6:].strip(), 11, SLATE)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:].strip(), 11, SLATE)
        elif line.startswith("|"):
            table_buf.append(line)
        else:
            text = line.strip()
            if text.startswith("**Como**") or text.startswith("**quero**") or text.startswith("**para**"):
                p = doc.add_paragraph()
                add_runs(p, text.replace("**", ""), 12, NAVY, True)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                add_runs(p, text.replace("**", ""), 11, SLATE)
        i += 1

    flush_table()
    flush_code()

    out_docs = DOCS / DOCX_NAME
    ART.mkdir(parents=True, exist_ok=True)
    out_art = ART / DOCX_NAME
    doc.save(str(out_docs))
    doc.save(str(out_art))
    print("Wrote", out_docs)
    print("Wrote", out_art)


if __name__ == "__main__":
    build()
